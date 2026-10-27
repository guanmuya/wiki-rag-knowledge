"""生成用于验证分块策略的测试文档：大表格 Word/Excel/PDF、多级标题 Word、纯文本。

用法：python scripts/make_fixtures.py
输出到 demo_data/fixtures/
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "demo_data" / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)

sys.path.insert(0, str(ROOT / "backend"))

COLS = ["工号", "姓名", "部门", "岗位", "入职日期", "职级", "月薪", "绩效等级", "备注"]
DEPTS = ["研发中心", "产品部", "市场部", "客户成功部", "财务部", "人力资源部"]
LEVELS = ["P4", "P5", "P6", "P7", "M1", "M2"]
GRADES = ["A", "B+", "B", "C"]


def gen_rows(n: int) -> list[list[str]]:
    rows = []
    for i in range(1, n + 1):
        rows.append([
            f"E{i:04d}",
            f"员工{i:03d}",
            DEPTS[i % len(DEPTS)],
            ["后端工程师", "前端工程师", "产品经理", "客户经理", "会计", "招聘专员"][i % 6],
            f"20{18 + i % 7}-{(i % 12) + 1:02d}-{(i % 27) + 1:02d}",
            LEVELS[i % len(LEVELS)],
            str(12000 + (i % 40) * 500),
            GRADES[i % len(GRADES)],
            "年度调薪已确认" if i % 3 == 0 else ("待复核" if i % 7 == 0 else ""),
        ])
    return rows


def make_docx_table(path: Path, rows: int) -> None:
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    doc.add_heading("员工薪酬明细表（2026 年度）", level=1)
    doc.add_paragraph(
        "本表汇总全公司员工的基础薪酬与职级信息，用于年度预算核算。"
        "数据来源于 HR 系统导出，口径为月度固定工资，不含奖金与补贴。")
    doc.add_heading("一、明细数据", level=2)

    data = gen_rows(rows)
    table = doc.add_table(rows=1, cols=len(COLS))
    table.style = "Table Grid"
    for i, name in enumerate(COLS):
        cell = table.rows[0].cells[i]
        cell.text = name
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
    for row in data:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)

    doc.add_heading("二、口径说明", level=2)
    doc.add_paragraph(
        "月薪为税前固定工资，绩效等级取最近一次季度考核结果。"
        "备注为空表示无需特殊说明。本表每季度更新一次，由人力资源部维护。")
    doc.save(str(path))


def make_xlsx(path: Path, rows: int) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "薪酬明细"
    ws.append(COLS)
    for row in gen_rows(rows):
        ws.append(row)
    ws2 = wb.create_sheet("部门汇总")
    ws2.append(["部门", "人数", "平均月薪", "预算占比"])
    for i, d in enumerate(DEPTS):
        ws2.append([d, str(rows // len(DEPTS) + (i % 3)), str(15000 + i * 800), f"{(i + 1) * 5}%"])
    wb.save(str(path))


def make_pdf(path: Path) -> None:
    """用 reportlab 不可用时退回纯文本 PDF（手写最小 PDF）。"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        c = canvas.Canvas(str(path), pagesize=A4)
        width, height = A4
        c.setFont("STSong-Light", 16)
        c.drawString(60, height - 60, "差旅与采购管理办法")
        c.setFont("STSong-Light", 10)
        lines = [
            "第一章 总则",
            "第一条 为规范公司差旅与采购行为，控制成本，制定本办法。",
            "第二条 本办法适用于公司全体员工的国内差旅与日常采购活动。",
            "",
            "第二章 差旅标准",
            "第三条 一线城市住宿标准为每人每晚 600 元，其他城市 400 元。",
            "第四条 交通工具标准为高铁二等座、飞机经济舱。",
            "第五条 出差补贴按每人每天 100 元计发，含餐费补助。",
            "",
            "第三章 采购流程",
            "第六条 单笔采购金额超过 5000 元的，需提交预算申请。",
            "第七条 采购需经部门负责人审批，财务复核后执行。",
            "第八条 供应商需在合格供应商名录内，名录每半年更新。",
            "",
            "第四章 附则",
            "第九条 本办法由财务部负责解释，自发布之日起施行。",
        ]
        y = height - 90
        for line in lines:
            if y < 60:
                c.showPage()
                c.setFont("STSong-Light", 10)
                y = height - 60
            c.drawString(60, y, line)
            y -= 18
        c.save()
    except ImportError:
        # 极简 PDF：仅验证流程时使用
        text = "差旅与采购管理办法\n\n一线城市住宿标准为每人每晚 600 元。\n"
        content = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET"
        objects = [
            "1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj",
            "2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj",
            "3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            "/Resources<< /Font<< /F1 4 0 R >> >> /Contents 5 0 R >>endobj",
            "4 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj",
            f"5 0 obj<< /Length {len(content)} >>stream\n{content}\nendstream endobj",
        ]
        pdf = "%PDF-1.4\n"
        offsets = []
        for obj in objects:
            offsets.append(len(pdf))
            pdf += obj + "\n"
        xref_at = len(pdf)
        pdf += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n"
        for off in offsets:
            pdf += f"{off:010d} 00000 n \n"
        pdf += (f"trailer<< /Size {len(objects)+1} /Root 1 0 R >>\n"
                f"startxref\n{xref_at}\n%%EOF\n")
        path.write_bytes(pdf.encode("latin-1"))


def make_multi_heading_docx(path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("研发流程规范", level=1)
    doc.add_paragraph("本规范描述需求、开发、测试与发布四个阶段的具体要求。")

    sections = [
        ("一、需求阶段", [
            "需求由产品经理统一收集，进入需求池后由评审会决定优先级。",
            "每个需求必须有明确的验收标准，缺失验收标准的需求不予排期。",
            "需求变更需评估影响范围，影响超过 5 人日需重新评审。",
        ]),
        ("二、开发阶段", [
            "开发任务以两周为迭代周期，任务拆分粒度不超过 3 人日。",
            "所有代码提交必须经过至少一人评审，主干分支禁止直接推送。",
            "单元测试覆盖率要求不低于 60%，核心模块不低于 80%。",
        ]),
        ("三、测试阶段", [
            "提测需提供自测报告，未通过自测的版本测试组有权拒收。",
            "回归测试范围包含核心链路与最近两个迭代的新功能。",
            "严重缺陷修复后需在测试环境验证通过才能进入发布流程。",
        ]),
        ("四、发布阶段", [
            "发布窗口为每周二、周四 20:00 至 23:00。",
            "发布前需准备回滚方案，并在灰度环境验证 30 分钟。",
            "发布完成后更新变更记录，异常情况在 24 小时内复盘。",
        ]),
    ]
    for title, paras in sections:
        doc.add_heading(title, level=2)
        for p in paras:
            doc.add_paragraph(p)
    doc.save(str(path))


def make_txt(path: Path) -> None:
    path.write_text(
        "客户投诉处理指引\n\n"
        "一、受理\n客户投诉可通过电话、邮件、工单三个渠道提交。"
        "受理人需在 2 小时内完成登记并分配处理人。\n\n"
        "二、处理\n处理人需在 1 个工作日内首次回复客户，"
        "并在 5 个工作日内给出处理结论。涉及退款的事项需财务会签。\n\n"
        "三、回访\n处理完成后 3 个工作日内完成回访，"
        "客户满意度低于 3 分的需升级至客服主管。\n",
        encoding="utf-8")


def main() -> None:
    make_docx_table(OUT / "员工薪酬明细_120行.docx", 120)
    make_docx_table(OUT / "员工薪酬明细_700行.docx", 700)
    make_xlsx(OUT / "员工薪酬明细_200行.xlsx", 200)
    make_pdf(OUT / "差旅与采购管理办法.pdf")
    make_multi_heading_docx(OUT / "研发流程规范.docx")
    make_txt(OUT / "客户投诉处理指引.txt")
    for f in sorted(OUT.iterdir()):
        print(f"  {f.name}  {f.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
